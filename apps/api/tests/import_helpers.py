"""Builders for the spreadsheet/PDF/Google fixtures used by the import tests."""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

HEADER: list[Any] = ["Customer Name", "Phone Number", "Amount Due", "Reference", "Due Date"]
GOOD: list[list[Any]] = [
    HEADER,
    ["Rahul Sharma", "9876543210", 15000, "INV-1001", date(2026, 10, 15)],
    ["Priya Traders", "98765 43211", "₹2,500.50", "INV-1002", "20/10/2026"],
    ["No Phone Co", None, 800, None, None],
]


def xlsx(rows: list[list[Any]], *, title_rows: int = 0, sheet: str | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    if sheet:
        ws.title = sheet
    for _ in range(title_rows):
        ws.append(["Monthly dues report"])
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def xlsx_sheets(sheets: dict[str, list[list[Any]]], hidden: tuple[str, ...] = ()) -> bytes:
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
        if name in hidden:
            ws.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def xls(rows: list[list[Any]]) -> bytes:
    import xlwt

    book = xlwt.Workbook()
    sh = book.add_sheet("Dues")
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if v is None:
                continue
            if isinstance(v, datetime | date):
                sh.write(r, c, v, date_style)
            else:
                sh.write(r, c, v)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def csv_bytes(text_: str, encoding: str = "utf-8") -> bytes:
    return text_.encode(encoding)


def pdf_table(rows: list[list[str]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    out = io.BytesIO()
    table = Table(rows)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(out, pagesize=A4).build([table])
    return out.getvalue()


def pdf_text(lines: list[str], pages: int = 1) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A4)
    for _ in range(pages):
        y = 800
        for line in lines:
            c.drawString(40, y, line)
            y -= 18
        c.showPage()
    c.save()
    return out.getvalue()


def scanned_pdf(lines: list[str], pages: int = 1) -> bytes:
    """A PDF that holds only a picture of text (no text layer), like a scan."""
    img = Image.new("RGB", (900, 600), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    y = 20
    for line in lines:
        draw.text((20, y), line, fill="black", font=font)
        y += 24
    out = io.BytesIO()
    img.save(out, format="PDF", save_all=True, append_images=[img.copy() for _ in range(pages - 1)])
    return out.getvalue()


def pdf_columns(rows: list[list[str]], pages: int = 1) -> bytes:
    """Text laid out in visible columns but with no ruling lines (so no table to extract)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A4)
    for _ in range(pages):
        y = 800
        for row in rows:
            for i, cell in enumerate(row):
                c.drawString(40 + i * 110, y, cell)
            y -= 18
        c.showPage()
    c.save()
    return out.getvalue()


def xlsx_merged(rows: list[list[Any]], merges: list[str], sheet: str = "Report") -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    for row in rows:
        ws.append(row)
    for ref in merges:
        ws.merge_cells(ref)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def tally_group_summary(parties: list[tuple[str, float | None, float | None]]) -> bytes:
    """A synthetic copy of the layout of a Tally 'Group Summary' export (no real data)."""
    rows: list[list[Any]] = [
        ["Example Traders Pvt Ltd, Pune"],
        ["1st Floor, Main Road, PUNE"],
        ["Pune, Maharashtra"],
        ["E-Mail : accounts@example.test"],
        ["Sundry Creditors"],
        ["Group Summary"],
        ["1-Apr-24 to 6-Jul-26"],
        [None, "Sundry Creditors"],
        [None, "Example Traders Pvt Ltd, Pune"],
        ["Particulars", "1-Apr-24 to 6-Jul-26"],
        [None, "Closing Balance"],
        [None, "Debit", "Credit"],
        *[[name, dr, cr] for name, dr, cr in parties],
        [
            "Grand Total",
            sum(p[1] or 0 for p in parties),
            sum(p[2] or 0 for p in parties),
        ],
    ]
    merges = [
        "A1:C1",
        "A2:C2",
        "A3:C3",
        "A4:C4",
        "A5:C5",
        "A6:C6",
        "A7:C7",
        "B8:C8",
        "B9:C9",
        "B10:C10",
        "B11:C11",
    ]
    return xlsx_merged(rows, merges, "Sundry Creditors")
