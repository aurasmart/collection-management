"""The petty-cash list as an Excel sheet (pure: rows in, .xlsx bytes out)."""

from __future__ import annotations

import io
from collections.abc import Sequence
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ["Date", "Transaction ID", "Payment to", "Payment from", "Remarks", "Amount (₹)"]
WIDTHS = [14, 24, 32, 32, 36, 16]
# Indian digit grouping (1,62,840.00), as used everywhere else in the app.
RUPEES = "[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00"
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def _text(value: str | None) -> str | None:
    """Receipt text is untrusted: stop a cell that starts like a formula from running as one."""
    if value and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


def build_xlsx(rows: Sequence[Any], total: Decimal) -> bytes:
    wb = Workbook()
    ws = wb.active or wb.create_sheet()
    ws.title = "Petty cash"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(
            [
                r.txn_date,
                _text(r.transaction_id),
                _text(r.payment_to),
                _text(r.payment_from),
                _text(r.remarks),
                r.amount,
            ]
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].number_format = "DD-MMM-YYYY"
        row[0].alignment = Alignment(horizontal="left")
        row[5].number_format = RUPEES
        for cell in row[1:5]:
            cell.data_type = "s"  # always text, never a formula
    ws.append(["", "", "", "", "Total", total])
    last = ws[ws.max_row]
    last[4].font = last[5].font = Font(bold=True)
    last[4].alignment = Alignment(horizontal="right")
    last[5].number_format = RUPEES
    for i, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
