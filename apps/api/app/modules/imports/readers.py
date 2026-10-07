"""File readers: bytes -> Workbook. No guessing about columns happens here."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook

from app.modules.imports.model import (
    MAX_COLUMNS,
    MAX_FILE_BYTES,
    MAX_ROWS,
    MAX_SHEETS,
    ImportFileError,
    Sheet,
    Workbook,
)

MAX_UNZIPPED_BYTES = 50 * 1024 * 1024
SCAN_ROWS = MAX_ROWS + 60  # data rows plus room for title/total/blank lines

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_SUPPORTED = "We can read Excel (.xlsx, .xls), CSV and PDF files. Please upload one of those."


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    """Drop trailing empty columns/rows and cap the width."""
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i] is not None and str(r[i]).strip() != "":
                width = max(width, i + 1)
                break
    out = [list(r[: min(width, MAX_COLUMNS)]) for r in rows]
    while out and not any(c is not None and str(c).strip() for c in out[-1]):
        out.pop()
    return out


def read_csv(data: bytes) -> Workbook:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    delimiter = max(",;\t|", key=first.count)
    try:
        rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)][:SCAN_ROWS]
    except csv.Error:
        raise ImportFileError(
            "We couldn't read this CSV file. Check that it is a valid CSV."
        ) from None
    return Workbook("csv", [Sheet("CSV", _trim(rows))])


def read_xlsx(data: bytes) -> Workbook:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise ImportFileError("This spreadsheet is too large to read.")
    except zipfile.BadZipFile:
        raise ImportFileError(
            "This file seems damaged. Try re-saving it and uploading again."
        ) from None
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise ImportFileError(
            "Spreadsheets with macros aren't supported. Save a copy as .xlsx without macros."
        )
    sheets: list[Sheet] = []
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets[:MAX_SHEETS]:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=SCAN_ROWS)]
            sheets.append(Sheet(ws.title, _trim(rows)))
        wb.close()
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError(
            "We couldn't read this spreadsheet. It may be damaged or password-protected."
        ) from None
    return Workbook("excel", _non_empty(sheets))


def read_xls(data: bytes) -> Workbook:
    import xlrd  # imported lazily: only legacy .xls files need it

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.XLRDError as exc:
        if "encrypt" in str(exc).lower() or "password" in str(exc).lower():
            raise ImportFileError("This file is password-protected. Remove the password.") from None
        raise ImportFileError(
            "We couldn't read this .xls file. Try saving it as .xlsx and upload again."
        ) from None
    except Exception:
        raise ImportFileError("This file seems damaged. Try re-saving it.") from None
    sheets: list[Sheet] = []
    try:
        for idx in range(min(book.nsheets, MAX_SHEETS)):
            sh = book.sheet_by_index(idx)
            if sh.visibility != 0:
                continue
            rows: list[list[Any]] = []
            for r in range(min(sh.nrows, SCAN_ROWS)):
                row: list[Any] = []
                for c in range(min(sh.ncols, MAX_COLUMNS + 1)):
                    row.append(_xls_cell(book, sh.cell(r, c)))
                rows.append(row)
            sheets.append(Sheet(sh.name, _trim(rows)))
    except xlrd.XLRDError:
        raise ImportFileError("We couldn't read this .xls file. Try saving it as .xlsx.") from None
    finally:
        book.release_resources()
    return Workbook("excel", _non_empty(sheets))


def _xls_cell(book: Any, cell: Any) -> Any:
    import xlrd

    t, v = cell.ctype, cell.value
    if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if t == xlrd.XL_CELL_DATE:
        try:
            value = xlrd.xldate.xldate_as_datetime(v, book.datemode)
        except Exception:
            return None
        return value
    if t == xlrd.XL_CELL_BOOLEAN:
        return bool(v)
    if t == xlrd.XL_CELL_NUMBER and float(v).is_integer():
        return int(v)
    return v


def _non_empty(sheets: list[Sheet]) -> list[Sheet]:
    kept = [s for s in sheets if s.rows]
    if not kept:
        raise ImportFileError("We couldn't find any data in this file.")
    return kept


def sniff(filename: str, data: bytes) -> str:
    """The real file type from its content first (extensions lie): xlsx, xls, csv or pdf."""
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("This file is larger than 5 MB. Split it and upload in parts.")
    if not data:
        raise ImportFileError("This file is empty.")
    name = (filename or "").lower()
    if data.startswith(b"PK\x03\x04"):
        if name.endswith((".docx", ".pptx", ".odt", ".zip")):
            raise ImportFileError(_SUPPORTED)
        return "xlsx"
    if data.startswith(_OLE_MAGIC):
        return "xls"
    if data.startswith(b"%PDF-"):
        return "pdf"
    if name.endswith((".xlsx", ".xls")):
        raise ImportFileError(
            "This doesn't look like a real Excel file. Open it in Excel and save it again."
        )
    if name.endswith(".pdf"):
        raise ImportFileError("This doesn't look like a real PDF file.")
    if name.endswith(".csv"):
        if b"\x00" in data[:4096]:
            raise ImportFileError("This doesn't look like a real CSV file.")
        return "csv"
    raise ImportFileError(_SUPPORTED)


def read_workbook(filename: str, data: bytes) -> Workbook:
    """Excel (.xlsx/.xls) and CSV. PDFs go through app.modules.imports.pdf (see sources.py)."""
    kind = sniff(filename, data)
    if kind == "xlsx":
        return read_xlsx(data)
    if kind == "xls":
        return read_xls(data)
    if kind == "csv":
        return read_csv(data)
    raise ImportFileError(_SUPPORTED)


# -- cell helpers shared with the detector ------------------------------------------------
def cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime) and v.time() == time(0, 0):
        return v.date().isoformat()  # an Excel date cell: show the date, not midnight
    if isinstance(v, datetime | date | time):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()
